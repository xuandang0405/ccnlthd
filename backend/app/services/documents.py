import re
import shutil
import unicodedata
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import BinaryIO

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader


class UnsafeDocumentError(ValueError):
    pass


MIME_BY_EXTENSION = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def save_limited(source: BinaryIO, destination: Path, maximum: int) -> int:
    destination.parent.mkdir(parents=True, exist_ok=True)
    written = 0
    try:
        with destination.open("xb") as target:
            while chunk := source.read(64 * 1024):
                written += len(chunk)
                if written > maximum:
                    raise UnsafeDocumentError(f"File vượt quá giới hạn {maximum} byte")
                target.write(chunk)
    except Exception:
        destination.unlink(missing_ok=True)
        raise
    if written == 0:
        destination.unlink(missing_ok=True)
        raise UnsafeDocumentError("File rỗng")
    return written


def _validate_zip(path: Path, required_member: str) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            if required_member not in names or "[Content_Types].xml" not in names:
                raise UnsafeDocumentError("Cấu trúc Office Open XML không hợp lệ")
            total_uncompressed = 0
            for item in archive.infolist():
                total_uncompressed += item.file_size
                if total_uncompressed > 100 * 1024 * 1024:
                    raise UnsafeDocumentError("Tài liệu nén có kích thước giải nén bất thường")
                if item.compress_size and item.file_size / item.compress_size > 200:
                    raise UnsafeDocumentError("Phát hiện tỷ lệ nén bất thường")
    except zipfile.BadZipFile as exc:
        raise UnsafeDocumentError("File Office bị hỏng hoặc không đúng định dạng") from exc


def validate_structure(path: Path, extension: str) -> str:
    extension = extension.lower()
    if extension not in MIME_BY_EXTENSION:
        raise UnsafeDocumentError("Chỉ hỗ trợ PDF, DOCX và XLSX")
    if extension == ".pdf":
        with path.open("rb") as stream:
            if stream.read(5) != b"%PDF-":
                raise UnsafeDocumentError("File không có cấu trúc PDF hợp lệ")
        try:
            PdfReader(path, strict=True)
        except Exception as exc:
            raise UnsafeDocumentError("File PDF bị hỏng hoặc không đọc được") from exc
    elif extension == ".docx":
        _validate_zip(path, "word/document.xml")
    else:
        _validate_zip(path, "xl/workbook.xml")
    return MIME_BY_EXTENSION[extension]


def normalize_text(value: str) -> str:
    value = value.replace("\x00", " ")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def normalize_search_text(value: str) -> str:
    value = unicodedata.normalize("NFKD", value.lower())
    value = "".join(character for character in value if not unicodedata.combining(character))
    return " ".join(re.findall(r"[a-z0-9]+", value))


def search_tokens(value: str, limit: int = 300) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for token in normalize_search_text(value).split():
        variants = [token]
        parts = re.findall(r"[a-z]+|\d+", token)
        if len(parts) > 1:
            variants.extend(parts)
        for variant in variants:
            if variant in seen or (len(variant) < 2 and not variant.isdigit()):
                continue
            seen.add(variant)
            result.append(variant)
            if len(result) >= limit:
                return result
    return result


def parse_pdf(path: Path, max_pages: int) -> str:
    reader = PdfReader(path, strict=True)
    if len(reader.pages) > max_pages:
        raise UnsafeDocumentError(f"PDF vượt quá giới hạn {max_pages} trang")
    return normalize_text("\n".join(page.extract_text() or "" for page in reader.pages))


def parse_docx(path: Path) -> str:
    document = DocxDocument(path)
    lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                lines.append(" | ".join(values))
    return normalize_text("\n".join(lines))


def _excel_value(value: object) -> str:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _xlsx_sheet_lines(sheet: object, max_cells: int, cells_seen: int) -> tuple[list[str], int]:
    rows = sheet.iter_rows(values_only=True)
    buffered: list[tuple[int, list[object]]] = []
    for row_number, row in enumerate(rows, start=1):
        values = list(row)
        nonempty = sum(value is not None and bool(str(value).strip()) for value in values)
        cells_seen += nonempty
        if cells_seen > max_cells:
            raise UnsafeDocumentError(f"XLSX vượt quá giới hạn xử lý {max_cells} ô có dữ liệu")
        if nonempty:
            buffered.append((row_number, values))
        if len(buffered) >= 20:
            break

    lines = [f"[Trang tính: {sheet.title}]"]
    if not buffered:
        return lines, cells_seen

    header_position = max(
        range(len(buffered)),
        key=lambda index: sum(
            value is not None and bool(str(value).strip()) for value in buffered[index][1]
        ),
    )
    header_row_number, header_values = buffered[header_position]
    headers = [
        _excel_value(value) if value is not None and str(value).strip() else f"Cột {index + 1}"
        for index, value in enumerate(header_values)
    ]

    for row_number, values in buffered[:header_position]:
        preamble = [
            _excel_value(value)
            for value in values
            if value is not None and str(value).strip()
        ]
        if preamble:
            lines.append(f"[Dòng {row_number}] " + " | ".join(preamble))

    def append_data_row(row_number: int, values: list[object]) -> None:
        fields = []
        for index, value in enumerate(values):
            if value is None or not str(value).strip():
                continue
            header = headers[index] if index < len(headers) else f"Cột {index + 1}"
            fields.append(f"{header}: {_excel_value(value)}")
        if fields:
            lines.append(
                f"[Trang tính: {sheet.title} | Dòng {row_number}] " + " | ".join(fields)
            )

    for row_number, values in buffered[header_position + 1 :]:
        append_data_row(row_number, values)

    next_row_number = buffered[-1][0] + 1
    for row_number, row in enumerate(rows, start=next_row_number):
        values = list(row)
        nonempty = sum(value is not None and bool(str(value).strip()) for value in values)
        cells_seen += nonempty
        if cells_seen > max_cells:
            raise UnsafeDocumentError(f"XLSX vượt quá giới hạn xử lý {max_cells} ô có dữ liệu")
        if nonempty:
            append_data_row(row_number, values)
    return lines, cells_seen


def parse_xlsx(path: Path, max_cells: int) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    lines: list[str] = []
    cells_seen = 0
    try:
        for sheet in workbook.worksheets:
            sheet_lines, cells_seen = _xlsx_sheet_lines(sheet, max_cells, cells_seen)
            lines.extend(sheet_lines)
    finally:
        workbook.close()
    return normalize_text("\n".join(lines))


def parse_document(path: Path, extension: str, max_pages: int, max_cells: int) -> str:
    if extension == ".pdf":
        text = parse_pdf(path, max_pages)
    elif extension == ".docx":
        text = parse_docx(path)
    elif extension == ".xlsx":
        text = parse_xlsx(path, max_cells)
    else:
        raise UnsafeDocumentError("Định dạng không được hỗ trợ")
    if not text:
        raise UnsafeDocumentError("Không trích xuất được nội dung văn bản")
    return text


def chunks(text: str, target_size: int = 2600, overlap: int = 240) -> list[str]:
    paragraphs = [part.strip() for part in text.split("\n") if part.strip()]
    result: list[str] = []
    current: list[str] = []
    for paragraph in paragraphs:
        current_size = sum(len(part) + 1 for part in current)
        if current and current_size + len(paragraph) + 1 > target_size:
            result.append("\n".join(current))
            tail: list[str] = []
            tail_size = 0
            for part in reversed(current):
                if tail and tail_size + len(part) + 1 > overlap:
                    break
                tail.insert(0, part)
                tail_size += len(part) + 1
            current = [*tail, paragraph]
        else:
            current.append(paragraph)
    if current:
        result.append("\n".join(current))
    return result


def safe_remove(path: Path) -> None:
    if path.is_file():
        path.unlink()
    elif path.exists():
        shutil.rmtree(path)
